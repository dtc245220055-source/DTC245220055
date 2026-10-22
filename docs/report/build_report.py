from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "report" / "Bao_cao_De_03_DTC245220055.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Times New Roman", size=13, bold=False, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    set_run_font(run, size=10)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 14, bold=True)
    return p


def add_body(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=True):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(6)
    if first_line:
        p.paragraph_format.first_line_indent = Cm(1)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_run_font(r, size=12.5)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "1F4E78")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, size=11, bold=True, color="FFFFFF")
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if ri % 2:
                set_cell_shading(cell, "EAF2F8")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, size=10.5)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_figure_slot(doc, label, caption, height_rows=4):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F2F2F2")
    set_cell_margins(cell, top=500, bottom=500, start=180, end=180)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(label)
    set_run_font(r, size=11, bold=True, color="666666")
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    rr = cap.add_run(caption)
    set_run_font(rr, size=10.5)
    rr.italic = True


def page_break(doc):
    doc.add_page_break()


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2)
section.bottom_margin = Cm(2)
section.left_margin = Cm(3)
section.right_margin = Cm(2)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal.font.size = Pt(13)
for style_name in ("Title", "Heading 1", "Heading 2"):
    style = doc.styles[style_name]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    style.font.color.rgb = RGBColor(0, 0, 0)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = header.add_run("Đề 03  |  DTC245220055")
set_run_font(hr, size=9, color="666666")
add_page_number(section.footer.paragraphs[0])

# Trang 1: Bìa
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(8)
r = p.add_run("TRƯỜNG ........................................................")
set_run_font(r, size=14, bold=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("KHOA ............................................................")
set_run_font(r, size=14, bold=True)
doc.add_paragraph("\n\n")
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("BÁO CÁO ĐỒ ÁN TRIỂN KHAI HỆ THỐNG")
set_run_font(r, size=20, bold=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("ĐỀ 03  WEBSITE QUẢNG BÁ NƯỚC RAU MÁ")
set_run_font(r, size=18, bold=True, color="1F4E78")
doc.add_paragraph("\n")
add_table(doc, ["Thông tin", "Nội dung"], [
    ("Sinh viên", "........................................................"),
    ("Mã sinh viên", "DTC245220055"),
    ("Lớp", "........................................................"),
    ("Giảng viên", "........................................................"),
    ("Môn học", "Triển khai và quản trị hệ thống"),
], widths=[4.2, 10.0])
doc.add_paragraph("\n")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Năm 2026")
set_run_font(r, size=13, bold=True)

# Trang 2: Mục lục
page_break(doc)
add_heading(doc, "MỤC LỤC", 1)
toc_rows = [
    ("1", "Giới thiệu đề tài", "3"), ("2", "Kiến trúc và công nghệ", "4"),
    ("3", "GitHub và cấu trúc dự án", "5"), ("4", "WordPress MySQL và phpMyAdmin", "6"),
    ("5", "Nginx Reverse Proxy", "7"), ("6", "Prometheus và Grafana", "8"),
    ("7", "Loki và Promtail", "9"), ("8", "Hardening hệ thống", "10"),
    ("9", "Kiểm thử và đánh giá", "11"), ("10", "Kịch bản bảo vệ và kết luận", "12"),
]
add_table(doc, ["Mục", "Nội dung", "Trang"], toc_rows, widths=[1.5, 10.5, 2.0])
add_body(doc, "Báo cáo trình bày toàn bộ quá trình triển khai theo bảy giai đoạn, tập trung vào kiến trúc container, khả năng quan sát hệ thống và các biện pháp giảm bề mặt tấn công. Các ảnh minh chứng được chọn ở mức tối thiểu để phù hợp giới hạn 10–15 trang.", first_line=False)

# Trang 3
page_break(doc)
add_heading(doc, "1 GIỚI THIỆU ĐỀ TÀI", 1)
add_body(doc, "Đề tài xây dựng website quảng bá sản phẩm nước rau má bằng WordPress và triển khai toàn bộ thành phần trên Docker Compose. Hệ thống không chỉ phục vụ nội dung website mà còn có lớp reverse proxy, cơ chế giám sát metrics, tập trung log và các cấu hình hardening cơ bản.")
add_heading(doc, "1.1 Mục tiêu", 2)
for x in [
    "Triển khai WordPress kết nối MySQL, hỗ trợ quản trị dữ liệu bằng phpMyAdmin.",
    "Đưa Nginx thành điểm truy cập duy nhất của website và bổ sung security headers.",
    "Giám sát dịch vụ, tài nguyên và cơ sở dữ liệu bằng Prometheus và Grafana.",
    "Thu thập log tập trung bằng Loki và Promtail, truy vấn bằng LogQL.",
    "Giảm rủi ro bằng cách cô lập mạng, hạn chế cổng công khai và tách tài khoản database.",
]: add_bullet(doc, x)
add_heading(doc, "1.2 Phạm vi và kết quả", 2)
add_body(doc, "Môi trường mục tiêu là Windows 11, Visual Studio Code và Docker Desktop. Hệ thống được xác minh trên máy cục bộ; không giả định máy ảo hoặc máy chủ từ xa. Kết quả cuối gồm mã nguồn cấu hình, ba mốc commit, tài liệu vận hành và kịch bản demo bảo vệ.")

# Trang 4
page_break(doc)
add_heading(doc, "2 KIẾN TRÚC VÀ CÔNG NGHỆ", 1)
add_table(doc, ["Lớp", "Thành phần", "Vai trò"], [
    ("Truy cập", "Nginx", "Reverse proxy, security headers, đầu vào duy nhất của website"),
    ("Ứng dụng", "WordPress", "Quản trị và hiển thị nội dung quảng bá sản phẩm"),
    ("Dữ liệu", "MySQL 8.4", "Lưu trữ dữ liệu WordPress trong named volume"),
    ("Quản trị", "phpMyAdmin", "Kiểm tra kết nối và cấu trúc database"),
    ("Metrics", "Prometheus, exporters", "Thu thập trạng thái Nginx, MySQL và container"),
    ("Hiển thị", "Grafana", "Dashboard metrics và khám phá log"),
    ("Logging", "Loki, Promtail", "Tập trung và truy vấn log container"),
], widths=[2.4, 4.2, 7.6])
add_heading(doc, "2.1 Luồng hoạt động", 2)
add_body(doc, "Người dùng truy cập localhost:8082. Nginx chuyển tiếp yêu cầu đến WordPress qua frontend network; WordPress kết nối MySQL qua backend network nội bộ. Các exporter cung cấp metrics cho Prometheus, còn Promtail đọc log container và đẩy đến Loki. Grafana sử dụng Prometheus và Loki làm hai nguồn dữ liệu.")
add_table(doc, ["Network", "Thuộc tính", "Nhóm dịch vụ"], [
    ("frontend", "internal=false", "Nginx, WordPress, phpMyAdmin"),
    ("backend", "internal=true", "WordPress, MySQL, phpMyAdmin, MySQL exporter"),
    ("monitoring", "internal=false", "Prometheus, Grafana, exporters, Loki, Promtail"),
], widths=[3.0, 3.5, 7.7])

# Trang 5
page_break(doc)
add_heading(doc, "3 GITHUB VÀ CẤU TRÚC DỰ ÁN", 1)
add_body(doc, "Repository DTC245220055 được tổ chức theo chức năng để dễ vận hành và kiểm tra. File .env chứa bí mật chỉ tồn tại cục bộ và đã được loại khỏi Git bằng .gitignore; repository chỉ lưu .env.example với giá trị mẫu.")
add_table(doc, ["Đường dẫn", "Nội dung"], [
    ("compose.yaml", "Khai báo toàn bộ service, volume, network và healthcheck"),
    ("nginx/", "Cấu hình reverse proxy và security headers"),
    ("monitoring/", "Prometheus, exporters, datasource và dashboard Grafana"),
    ("logging/", "Cấu hình Loki, Promtail và provisioning datasource"),
    ("docs/evidence/", "Ảnh minh chứng dùng trong báo cáo"),
    ("docs/report/", "Báo cáo và tài liệu bảo vệ"),
], widths=[4.5, 9.7])
add_heading(doc, "3.1 Ba mốc commit", 2)
add_table(doc, ["Mốc", "Commit", "Phạm vi"], [
    ("1", "f0514fa", "WordPress, MySQL, phpMyAdmin và Nginx"),
    ("2", "03be795", "Prometheus, Grafana và exporters"),
    ("3", "e38b8ac", "Loki, Promtail và hardening"),
], widths=[1.5, 3.0, 9.7])
add_figure_slot(doc, "ẢNH GITHUB REPOSITORY VÀ LỊCH SỬ 3 COMMIT", "Hình 3.1. Repository GitHub của dự án và ba mốc commit theo yêu cầu.")

# Trang 6
page_break(doc)
add_heading(doc, "4 WORDPRESS MYSQL VÀ PHPMYADMIN", 1)
add_body(doc, "WordPress sử dụng image wordpress:php8.3-apache và kết nối đến MySQL 8.4 bằng biến môi trường. MySQL có healthcheck và lưu dữ liệu trong named volume, giúp dữ liệu không mất khi container được tạo lại. phpMyAdmin chỉ bind trên địa chỉ loopback để hỗ trợ quản trị cục bộ.")
add_table(doc, ["Dịch vụ", "Cổng host", "Kết quả"], [
    ("WordPress", "Không publish trực tiếp", "Truy cập qua Nginx"),
    ("MySQL", "Không publish 3306", "Healthy, chỉ dùng backend network"),
    ("phpMyAdmin", "127.0.0.1:8083", "Kết nối được MySQL"),
], widths=[4.0, 4.2, 6.0])
add_figure_slot(doc, "ẢNH DOCKER COMPOSE PS VÀ MYSQL HEALTHY", "Hình 4.1. Các container WordPress, MySQL và phpMyAdmin hoạt động; MySQL ở trạng thái healthy.")
add_figure_slot(doc, "ẢNH TRANG CHỦ WORDPRESS HOẶC PHPMYADMIN", "Hình 4.2. Website WordPress và kết nối cơ sở dữ liệu được kiểm tra thành công.")

# Trang 7
page_break(doc)
add_heading(doc, "5 NGINX REVERSE PROXY", 1)
add_body(doc, "Nginx là điểm vào duy nhất của website tại http://localhost:8082. Cấu hình proxy chuyển tiếp request đến wordpress:80 bằng tên service trong Docker network. WordPress không còn publish cổng trực tiếp ra host, nhờ đó tránh truy cập vòng qua lớp proxy.")
add_heading(doc, "5.1 Security headers", 2)
add_table(doc, ["Header", "Giá trị"], [
    ("X-Frame-Options", "SAMEORIGIN"),
    ("X-Content-Type-Options", "nosniff"),
    ("Referrer-Policy", "strict-origin-when-cross-origin"),
    ("Permissions-Policy", "camera=(), microphone=(), geolocation=()"),
], widths=[5.5, 8.7])
add_body(doc, "Lệnh curl.exe -I http://localhost:8082 trả về HTTP/1.1 200 OK, Server: nginx và đầy đủ bốn header bảo vệ. Cấu hình cũng tắt server_tokens để giảm thông tin phiên bản lộ ra ngoài.")
add_figure_slot(doc, "ẢNH TERMINAL HTTP 200 VÀ SECURITY HEADERS", "Hình 5.1. Nginx trả về HTTP 200 và các security headers đã cấu hình.")

# Trang 8
page_break(doc)
add_heading(doc, "6 PROMETHEUS VÀ GRAFANA", 1)
add_body(doc, "Prometheus thu thập metrics theo mô hình pull. Bốn target gồm Prometheus, cAdvisor, Nginx Exporter và MySQL Exporter đều được kiểm tra ở trạng thái UP. Grafana được provision tự động datasource Prometheus và dashboard De 03 WordPress Monitoring.")
add_table(doc, ["Target", "Dữ liệu chính", "Trạng thái"], [
    ("prometheus", "Tình trạng máy chủ metrics", "UP"),
    ("cadvisor", "CPU, RAM, network của container", "UP"),
    ("nginx", "Kết nối và request Nginx", "UP"),
    ("mysql", "Kết nối và hoạt động MySQL", "UP"),
], widths=[3.2, 7.8, 3.2])
add_figure_slot(doc, "ẢNH PROMETHEUS TARGETS", "Hình 6.1. Prometheus thu thập thành công metrics từ bốn targets; tất cả ở trạng thái UP.")
add_figure_slot(doc, "ẢNH GRAFANA DASHBOARD", "Hình 6.2. Dashboard Grafana giám sát trạng thái dịch vụ và tài nguyên container.")

# Trang 9
page_break(doc)
add_heading(doc, "7 LOKI VÀ PROMTAIL", 1)
add_body(doc, "Promtail tự động phát hiện container thuộc Compose project, gắn nhãn service và gửi log đến Loki. Loki lưu dữ liệu trong named volume; Grafana đã được provision datasource Loki tại địa chỉ nội bộ http://loki:3100.")
add_heading(doc, "7.1 Các truy vấn LogQL", 2)
add_table(doc, ["Mục đích", "Truy vấn"], [
    ("Xem toàn bộ log Nginx", '{service="nginx"}'),
    ("Lọc lỗi HTTP 4xx", '{service="nginx"} |~ " 4[0-9]{2} "'),
    ("Lọc request GET WordPress", '{service="wordpress"} |= "GET"'),
], widths=[5.0, 9.2])
add_figure_slot(doc, "ẢNH LOGQL LOG NGINX", "Hình 7.1. Grafana Explore truy vấn thành công log Nginx do Promtail gửi đến Loki.")
add_figure_slot(doc, "ẢNH LOGQL LỌC 4XX HOẶC 404", "Hình 7.2. Truy vấn LogQL lọc yêu cầu HTTP lỗi 4xx trong log Nginx.")

# Trang 10
page_break(doc)
add_heading(doc, "8 HARDENING HỆ THỐNG", 1)
add_body(doc, "Hardening tập trung vào quyền container, phân đoạn mạng, bí mật, quyền database và các cổng được công khai. Việc kiểm tra được thực hiện bằng docker inspect, docker network inspect, docker compose ps và câu lệnh SHOW GRANTS.")
add_table(doc, ["Hạng mục", "Kết quả kiểm tra"], [
    ("Privileged", "Chỉ cAdvisor=true để đọc metrics; các container còn lại=false"),
    ("Backend network", "internal=true, cô lập MySQL khỏi truy cập ngoài"),
    ("Cổng công khai", "Chỉ bind 127.0.0.1; MySQL và WordPress không publish trực tiếp"),
    ("Bí mật", ".env bị gitignore; mật khẩu dài 24–29 ký tự, không dùng placeholder"),
    ("Database", "wp_rauma chỉ có quyền trên wordpress_rauma.*"),
    ("Grafana", "Tắt tự đăng ký tài khoản"),
    ("Mount", "Cấu hình và Docker socket phục vụ giám sát được mount read-only"),
], widths=[4.0, 10.2])
add_figure_slot(doc, "ẢNH TERMINAL SHOW GRANTS", "Hình 8.1. Tài khoản wp_rauma chỉ được cấp quyền trên cơ sở dữ liệu wordpress_rauma, không dùng MySQL root.")

# Trang 11
page_break(doc)
add_heading(doc, "9 KIỂM THỬ VÀ ĐÁNH GIÁ", 1)
add_table(doc, ["Kiểm thử", "Cách kiểm tra", "Kết quả"], [
    ("Container", "docker compose ps", "Các dịch vụ chạy; healthcheck đạt yêu cầu"),
    ("Web", "curl -I localhost:8082", "HTTP 200 qua Nginx"),
    ("Database", "phpMyAdmin và SHOW GRANTS", "Kết nối đúng, tài khoản tách biệt"),
    ("Metrics", "Prometheus /targets", "4 targets UP"),
    ("Dashboard", "Grafana", "Dashboard tải dữ liệu thành công"),
    ("Logging", "Grafana Explore", "Log Nginx/WordPress truy vấn được"),
    ("Git", "git log và git status", "Đúng 3 commit, không track .env"),
], widths=[3.0, 5.5, 5.7])
add_heading(doc, "9.1 Đánh giá", 2)
add_body(doc, "Hệ thống đáp ứng đầy đủ các lớp chức năng của đề bài trong môi trường cục bộ. Cấu hình có khả năng tái tạo bằng một lệnh Docker Compose, dữ liệu quan trọng được lưu trong volume và dashboard/datasource được provision tự động. Bề mặt tấn công được giảm bằng cách không công khai MySQL và WordPress, cô lập backend network và chỉ bind giao diện quản trị vào loopback.")
add_heading(doc, "9.2 Hạn chế và hướng phát triển", 2)
add_body(doc, "Môi trường hiện tại dùng HTTP trên localhost để phục vụ demo. Khi triển khai Internet cần tên miền, chứng chỉ TLS hợp lệ, backup định kỳ, giới hạn tài nguyên, chính sách cập nhật image và cơ chế cảnh báo. phpMyAdmin không nên mở công khai trong môi trường production.")

# Trang 12
page_break(doc)
add_heading(doc, "10 KỊCH BẢN BẢO VỆ VÀ KẾT LUẬN", 1)
add_heading(doc, "10.1 Kịch bản demo 7 phút", 2)
add_table(doc, ["Thời gian", "Nội dung trình bày"], [
    ("0:00–0:45", "Giới thiệu mục tiêu, stack và sơ đồ kiến trúc"),
    ("0:45–1:30", "Mở GitHub, cấu trúc thư mục và lịch sử đúng 3 commit"),
    ("1:30–2:30", "docker compose ps; mở WordPress và phpMyAdmin"),
    ("2:30–3:15", "curl -I qua Nginx, chỉ ra HTTP 200 và security headers"),
    ("3:15–4:30", "Prometheus targets UP và dashboard Grafana"),
    ("4:30–5:45", "Grafana Explore, chạy LogQL tổng quát và lọc 4xx"),
    ("5:45–6:40", "Trình bày network internal, port binding và SHOW GRANTS"),
    ("6:40–7:00", "Kết luận và hướng phát triển HTTPS production"),
], widths=[3.0, 11.2])
add_heading(doc, "10.2 Câu hỏi bảo vệ dự kiến", 2)
for x in [
    "Vì sao không publish cổng WordPress? Để mọi truy cập đi qua Nginx và áp dụng chung chính sách bảo vệ.",
    "Vì sao MySQL không mở cổng 3306? Ứng dụng kết nối qua backend network nên không cần công khai ra host.",
    "Prometheus khác Loki thế nào? Prometheus lưu metrics dạng chuỗi thời gian; Loki lưu và truy vấn log.",
    "Tại sao cAdvisor privileged? Nó cần quyền đọc thông tin container và host; đây là ngoại lệ đã được ghi nhận.",
]: add_bullet(doc, x)
add_heading(doc, "10.3 Kết luận", 2)
add_body(doc, "Đồ án đã triển khai thành công website WordPress cùng cơ sở dữ liệu, reverse proxy, metrics, dashboard, logging tập trung và hardening cơ bản. Ba mốc commit thể hiện rõ tiến trình xây dựng. Hệ thống có thể tái tạo trên máy cài Docker Desktop và sẵn sàng cho phần trình diễn bảo vệ.")

doc.core_properties.title = "Báo cáo Đề 03 Website quảng bá nước rau má"
doc.core_properties.subject = "Triển khai WordPress Docker Compose Nginx Prometheus Grafana Loki Promtail"
doc.core_properties.author = "DTC245220055"
doc.core_properties.keywords = "WordPress, Docker, Nginx, Prometheus, Grafana, Loki, Promtail"
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print("report-created")
